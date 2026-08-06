"""
ITM-style Student Attendance Register Excel export.

Matches the paper register concept:
  - Header: Session, Class/Branch, Subject & Code, Faculty
  - Grid: S.No | Roll No | Name | date columns with P
  - Summary: Total Attnd | Total Attnd % | blank columns for
    Regularity, Performance, Internal Marks, Grand Total, Grade Point

Teachers can regenerate this file any day (always up to date) and
submit the monthly sheet to higher authority.
"""

from __future__ import annotations

import calendar
import datetime
import io
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import db


def _month_days(year: int, month: int) -> list[datetime.date]:
    n = calendar.monthrange(year, month)[1]
    return [datetime.date(year, month, d) for d in range(1, n + 1)]


def _thin():
    return Border(
        left=Side(style="thin", color="999999"),
        right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"),
        bottom=Side(style="thin", color="999999"),
    )


def build_register_workbook(
    *,
    class_id: int,
    subject_id: int,
    year: int,
    month: int,
    session_label: str = "",
    faculty_name: str = "",
    branch_label: str = "",
) -> bytes:
    """
    Return .xlsx bytes for one class+subject+month register sheet.
    """
    with db.get_conn() as conn:
        crow = conn.execute(
            "SELECT id, name, section, academic_year FROM classes WHERE id=?",
            (class_id,),
        ).fetchone()
        srow = conn.execute(
            "SELECT id, name, code FROM subjects WHERE id=? AND class_id=?",
            (subject_id, class_id),
        ).fetchone()
        if not crow or not srow:
            raise ValueError("Class or subject not found")

        students = conn.execute(
            """
            SELECT DISTINCT st.id, st.name, st.roll
            FROM students st
            LEFT JOIN enrollments e ON e.student_id = st.id
            WHERE st.class_id=? OR e.class_id=?
            ORDER BY st.roll, st.name
            """,
            (class_id, class_id),
        ).fetchall()

        # Present set: (student_id, YYYY-MM-DD)
        start = f"{year:04d}-{month:02d}-01"
        if month == 12:
            end = f"{year + 1:04d}-01-01"
        else:
            end = f"{year:04d}-{month + 1:02d}-01"

        marks = conn.execute(
            """
            SELECT student_id, attendance_day
            FROM attendance
            WHERE class_id=? AND subject_id=?
              AND attendance_day >= ? AND attendance_day < ?
            """,
            (class_id, subject_id, start, end),
        ).fetchall()

    present = {(int(r["student_id"]), r["attendance_day"]) for r in marks if r["attendance_day"]}
    days = _month_days(year, month)

    class_name = crow["name"] or ""
    section = crow["section"] or ""
    class_display = class_name + (f" — Sec {section}" if section else "")
    subject_name = srow["name"] or ""
    subject_code = srow["code"] or ""
    session_label = session_label or (crow["academic_year"] or f"{year}-{str(year + 1)[-2:]}")
    branch_label = branch_label or class_name

    wb = Workbook()

    # ----- Cover-like info sheet -----
    cover = wb.active
    cover.title = "Register Cover"
    cover["A1"] = "ITM UNIVERSITY — DIGITAL ATTENDANCE REGISTER"
    cover["A1"].font = Font(bold=True, size=16)
    cover.merge_cells("A1:D1")

    cover["A3"] = "Student Attendance (Theory / Practical)"
    cover["A3"].font = Font(bold=True, size=13)

    info = [
        ("Session", session_label),
        ("Class", class_display),
        ("Branch", branch_label),
        ("Subject & Code", f"{subject_name}" + (f" ({subject_code})" if subject_code else "")),
        ("Faculty Incharge", faculty_name or "—"),
        ("Month", f"{calendar.month_name[month]} {year}"),
        ("Generated on", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(info, start=5):
        cover[f"A{i}"] = k
        cover[f"A{i}"].font = Font(bold=True)
        cover[f"B{i}"] = v
    cover.column_dimensions["A"].width = 22
    cover.column_dimensions["B"].width = 48
    cover["A13"] = (
        "Instruction: After each day's attendance is marked in the app, re-download this "
        "monthly register. At month end, print/sign and submit to higher authority "
        "(Dean / HoD). Columns for Regularity, Internal Marks and Grade can be filled manually."
    )
    cover.merge_cells("A13:D16")
    cover["A13"].alignment = Alignment(wrap_text=True, vertical="top")

    # ----- Main grid sheet -----
    ws = wb.create_sheet("Monthly Register")
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    # Title block
    ws.merge_cells("A1:H1")
    ws["A1"] = "Student Attendance Register (Digital)"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = f"Branch & Semester / Class: {class_display}"
    ws["A3"] = f"Subject Name: {subject_name}"
    ws["D3"] = f"Subject Code: {subject_code or '—'}"
    ws["A4"] = f"Faculty Incharge: {faculty_name or '—'}"
    ws["D4"] = f"Session: {session_label} | Month: {calendar.month_name[month]} {year}"

    # Column headers start at row 6
    # Fixed cols: S.No, Roll No, Name | then each day | totals + extra paper-register cols
    fixed = ["S.No.", "Roll No.", "Name of the Student"]
    extra = [
        "Total Attnd.",
        "Total Attnd. %",
        "Regularity/Conduct",
        "Student Performance",
        "Internal Marks I",
        "Internal Marks II",
        "Internal Marks III",
        "Grand Total",
        "Grade Point",
    ]

    header_row = 6
    date_row = 7  # short date labels under day index optional — put dates in header

    # Row 6: S.No Roll Name + day numbers 1..n + extras
    # Row 7: blank/blank/blank + dd/mm labels + blank extras titles already on row 6

    for col, title in enumerate(fixed, start=1):
        cell = ws.cell(header_row, col, title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = center
        cell.border = _thin()

    for i, day in enumerate(days):
        col = 3 + i + 1
        # Day index like paper register
        c1 = ws.cell(header_row, col, i + 1)
        c1.font = Font(bold=True, size=9)
        c1.fill = header_fill
        c1.alignment = center
        c1.border = _thin()
        c2 = ws.cell(date_row, col, day.strftime("%d/%m"))
        c2.font = Font(size=8)
        c2.alignment = Alignment(horizontal="center", textRotation=90, wrap_text=True)
        c2.border = _thin()

    first_extra = 3 + len(days) + 1
    for j, title in enumerate(extra):
        col = first_extra + j
        cell = ws.cell(header_row, col, title)
        cell.font = Font(bold=True, size=9)
        cell.fill = header_fill
        cell.alignment = center
        cell.border = _thin()
        ws.cell(date_row, col, "").border = _thin()

    # Fill fixed header date row blanks
    for col in range(1, 4):
        ws.cell(date_row, col, "").border = _thin()
        ws.cell(date_row, col).fill = header_fill

    ws.row_dimensions[date_row].height = 45

    working_days = len(days)  # percentage base = days in month (college may redefine)

    data_start = 8
    for idx, st in enumerate(students, start=1):
        row = data_start + idx - 1
        sid = int(st["id"])
        roll = st["roll"] or ""
        name = st["name"] or ""

        values = [idx, roll, name]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row, col, val)
            cell.border = _thin()
            cell.alignment = left if col == 3 else center

        attended = 0
        for i, day in enumerate(days):
            col = 3 + i + 1
            key = (sid, day.isoformat())
            mark = "P" if key in present else ""
            if mark:
                attended += 1
            cell = ws.cell(row, col, mark)
            cell.alignment = center
            cell.border = _thin()
            if mark:
                cell.font = Font(bold=True, color="006600")

        pct = round((attended / working_days) * 100, 1) if working_days else 0.0
        summary_vals = [attended, pct, "", "", "", "", "", "", ""]
        for j, val in enumerate(summary_vals):
            col = first_extra + j
            cell = ws.cell(row, col, val)
            cell.border = _thin()
            cell.alignment = center

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 28
    for i in range(len(days)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 3.2
    for j in range(len(extra)):
        ws.column_dimensions[get_column_letter(first_extra + j)].width = 14

    # Signature block
    sig_row = data_start + len(students) + 2
    ws.cell(sig_row, 1, "Faculty Signature: ____________________")
    ws.cell(sig_row, first_extra, "Dean / HoD: ____________________")
    ws.cell(sig_row + 1, 1, "Note: P = Present. Blank = Absent / Not marked. Re-download after each day to keep Excel updated.")

    # ----- Daily log sheet (optional helpful) -----
    log = wb.create_sheet("Daily Log")
    log["A1"] = "Daily attendance log (same month)"
    log["A1"].font = Font(bold=True)
    for col, h in enumerate(["Date", "Roll No.", "Name", "Source"], start=1):
        cell = log.cell(3, col, h)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = _thin()

    # Fetch detailed rows
    with db.get_conn() as conn:
        detail = conn.execute(
            """
            SELECT a.attendance_day, a.name, a.source, st.roll
            FROM attendance a
            LEFT JOIN students st ON st.id = a.student_id
            WHERE a.class_id=? AND a.subject_id=?
              AND a.attendance_day >= ? AND a.attendance_day < ?
            ORDER BY a.attendance_day, st.roll, a.name
            """,
            (class_id, subject_id, start, end),
        ).fetchall()

    r = 4
    for row in detail:
        day = row["attendance_day"] or ""
        try:
            day_fmt = datetime.date.fromisoformat(day).strftime("%d/%m/%Y")
        except Exception:
            day_fmt = day
        vals = [day_fmt, row["roll"] or "", row["name"] or "", row["source"] or ""]
        for c, v in enumerate(vals, start=1):
            cell = log.cell(r, c, v)
            cell.border = _thin()
        r += 1
    for col in range(1, 5):
        log.column_dimensions[get_column_letter(col)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
